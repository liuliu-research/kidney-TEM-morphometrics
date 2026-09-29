"""抽取字段级人工核对样本：复现冻结建模队列，分层随机抽样，生成盲法读取表与密钥文件。

只读取 0病理数据清洗与分类 的锁定注册表、旧形态表与原始报告；不改动任何源数据。
输出：
  audit_key_CONFIDENTIAL.json  —— 样本与机器抽取值的对照密钥（不给读者看，不入公开仓库）
  audit_worksheet_reader1.xlsx —— 盲法读取表（读者 1）
  audit_worksheet_reader2.xlsx —— 盲法读取表（读者 2，可选）
  audit_sample_manifest.json   —— 抽样参数、哈希与队列复现核对
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd

DATA_ROOT = Path(os.environ["PATHOLOGY_V2_ROOT"]).expanduser().resolve()
os.environ["PATHOLOGY_V2_ROOT"] = str(DATA_ROOT)
PROJECT_ROOT = Path(os.environ.get("TEM_FROZEN_REPO", str(Path(__file__).resolve().parents[2]))).resolve()
REPO_SCRIPTS = PROJECT_ROOT / "scripts"
FROZEN_RESULTS = (
    PROJECT_ROOT / "aggregate_results" / "v2_modeling_results.json"
)
OUT_DIR = Path(os.environ["TEM_AUDIT_OUTPUT_DIR"]).expanduser().resolve()
SEED = 20260929
PER_STRATUM = 100
STRATA = ["development", "temporal_test"]

sys.path.insert(0, str(REPO_SCRIPTS))
import remote_v2_modeling as rvm  # noqa: E402

FIELD_SPECS = [
    ("glomerulus_count", "肾小球总数", "肾小球总数（个）"),
    ("sclerosis_rate", "肾小球硬化", "肾小球硬化（球性 N 个，节段性 M 个）"),
    ("mesangial_cell_count", "系膜细胞数/区", "系膜细胞数/区（如 1-2、3-6）"),
    ("mesangial_matrix_grade", "系膜基质增生程度", "系膜基质增生程度（无/轻度/轻-中度/中度/中-重度/重度）"),
    ("gbm_thickness_nm", "基底膜厚度", "基底膜厚度（如 300-500nm）"),
    ("spike_presence", "钉突有无", "钉突（形成/未见/不明确）"),
    ("foot_process_grade", "足突融合程度", "足突融合程度（无/节段性或部分/较广泛或大部分/广泛性）"),
]


def sha256_file(path: Path) -> str:
    """返回文件 SHA-256。"""

    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def resolve_report(report_path: str, stem: str, visit_id: str) -> Path:
    """定位 v2.0 数据地基中的匿名报告：文件已按 <visit_id>.txt 重命名，注册表旧路径仅作回退。"""

    hits = list((DATA_ROOT / "肾病分类").glob(f"*/诊断报告/{visit_id}.txt"))
    if len(hits) == 1:
        return hits[0]
    direct = DATA_ROOT / report_path
    if direct.is_file():
        return direct
    hits = list((DATA_ROOT / "肾病分类").glob(f"*/诊断报告/{stem}.txt"))
    if len(hits) != 1:
        raise SystemExit(f"[FAIL] 报告无法唯一定位: visit_id={visit_id!r} stem={stem!r} -> {len(hits)} 个")
    return hits[0]


def report_excerpt(text: str) -> str:
    """只保留电镜观察描述与电镜诊断，去掉抬头行（姓名/住院号/床号）。"""

    lines = [line.rstrip() for line in unicodedata.normalize("NFC", text).splitlines()]
    keep: list[str] = []
    started = False
    for line in lines:
        stripped = line.strip()
        if not started:
            if stripped.startswith("电镜") and ("观察" in stripped or "描述" in stripped):
                started = True
            elif stripped.startswith("肾小球："):
                started = True
        if started:
            if stripped.startswith("签名") or stripped.startswith("诊断医生") or stripped == "电镜图片":
                break
            if stripped:
                keep.append(stripped)
    if not keep:
        raise SystemExit("[FAIL] 报告缺少可识别的电镜观察段落")
    return "\n".join(keep)


def main() -> None:
    """生成样本、读取表与密钥。"""

    linked, audit = rvm.build_dataset()
    frozen = json.loads(FROZEN_RESULTS.read_text(encoding="utf-8"))["audit"]
    checks = {
        "counts_match_frozen": audit["counts"] == frozen["counts"],
        "linkage_match_frozen": audit["linkage"]["method_counts"] == frozen["linkage"]["method_counts"],
        "missingness_match_frozen": audit["feature_missingness_pct"] == frozen["feature_missingness_pct"],
        "class_counts_match_frozen": audit["class_counts"] == frozen["class_counts"],
    }
    if not all(checks.values()):
        raise SystemExit(f"[FAIL] 本地复现与冻结结果不一致: {checks}")

    legacy = pd.read_excel(rvm.MORPHOMETRY, sheet_name="分类汇总")
    rng = np.random.default_rng(SEED)
    rows: list[dict[str, object]] = []
    for stratum in STRATA:
        pool = linked.loc[linked["split"] == stratum].sort_values(["date", "visit_id"], kind="stable")
        chosen = rng.choice(pool.index.to_numpy(), size=PER_STRATUM, replace=False)
        for idx in sorted(int(i) for i in chosen):
            rec = linked.loc[idx]
            legacy_row = legacy.loc[int(rec["legacy_index"])]
            path = resolve_report(str(rec["report_path"]), str(rec["report_stem"]), str(rec["visit_id"]))
            text = path.read_text(encoding="utf-8", errors="replace")
            rows.append(
                {
                    "linked_index": idx,
                    "split": stratum,
                    "PT": rec["PT"],
                    "visit_id": rec["visit_id"],
                    "song_ymd": str(rec["song_ymd"]),
                    "disease_code": rec["disease_code"],
                    "link_method": rec["link_method"],
                    "legacy_index": int(rec["legacy_index"]),
                    "report_file": str(path.relative_to(DATA_ROOT)),
                    "report_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    "excerpt": report_excerpt(text),
                    "raw_cells": {col: (None if pd.isna(legacy_row[col]) else str(legacy_row[col])) for _, col, _ in FIELD_SPECS},
                    "parsed": {
                        key: (None if pd.isna(rec[key]) else float(rec[key])) for key, _, _ in FIELD_SPECS
                    },
                }
            )
    # 随机打乱样本顺序并赋匿名样本号，读者看不到分期与类别
    order = rng.permutation(len(rows))
    for rank, position in enumerate(order, 1):
        rows[int(position)]["sample_id"] = f"S{rank:03d}"
    rows.sort(key=lambda r: r["sample_id"])

    key_path = OUT_DIR / "audit_key_CONFIDENTIAL.json"
    key_path.write_text(
        json.dumps(
            {
                "seed": SEED,
                "per_stratum": PER_STRATUM,
                "strata": STRATA,
                "fields": [{"key": k, "source_column": c, "label": l} for k, c, l in FIELD_SPECS],
                "rows": rows,
            },
            ensure_ascii=False,
            indent=1,
        ),
        encoding="utf-8",
    )

    guide = pd.DataFrame(
        {
            "项目": [
                "目的",
                "盲法",
                "读取规则",
                "肾小球总数",
                "肾小球硬化",
                "系膜细胞数/区",
                "系膜基质增生程度",
                "基底膜厚度",
                "钉突",
                "足突融合程度",
                "空白",
                "备注",
            ],
            "说明": [
                "对随机抽取的 200 份电镜报告，按报告原文逐字段抄录，用于与机器抽取值比对并报告每个字段的错误率。",
                "读取表不含机器抽取值、分期或分类；请勿查看密钥文件；两位读者独立完成，不互相讨论。",
                "只抄录报告明确写出的内容；报告未提及或措辞模糊（如“可能”“不明显”）时填写“未提及”或“不明确”，不要推断。",
                "填写报告写明的肾小球总数（整数）；报告给范围时按原文填写，如“3-5”。",
                "按原文填写球性与节段性硬化个数，如“球性1个，节段性2个”；“趋于/接近/缺血性球性硬化”不算明确计数，按原文抄录并在备注注明。",
                "按原文填写，如“1-2”“3-6”“4”。",
                "从以下选项选择：无增生 / 轻度 / 轻-中度 / 中度 / 中-重度 / 重度 / 未提及 / 不明确。只看“基质”增生，不看细胞增生。",
                "按原文填写厚度及单位，如“500-1000nm”“0.3-0.5μm”；无数值则填“未提及”。",
                "从以下选项选择：形成 / 未见 / 不明确 / 未提及。",
                "从以下选项选择：无融合 / 节段性或部分或少数 / 较广泛或大部分 / 广泛性 / 未提及 / 不明确。",
                "未提及 = 报告根本没写该字段；不明确 = 写了但措辞模糊无法归类。",
                "任何疑问写在备注列。完成后保存并交回，勿改动样本号与报告原文。",
            ],
        }
    )
    sheet = pd.DataFrame(
        {
            "样本号": [r["sample_id"] for r in rows],
            "报告原文（电镜观察与诊断）": [r["excerpt"] for r in rows],
            "肾小球总数": "",
            "肾小球硬化": "",
            "系膜细胞数/区": "",
            "系膜基质增生程度": "",
            "基底膜厚度": "",
            "钉突": "",
            "足突融合程度": "",
            "备注": "",
        }
    )
    for reader in (1, 2):
        out = OUT_DIR / f"audit_worksheet_reader{reader}.xlsx"
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            guide.to_excel(writer, sheet_name="说明", index=False)
            sheet.to_excel(writer, sheet_name="读取表", index=False)
            ws = writer.sheets["读取表"]
            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 110
            for col in "CDEFGHIJ":
                ws.column_dimensions[col].width = 18
            from openpyxl.styles import Alignment
            from openpyxl.worksheet.datavalidation import DataValidation

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                row[1].alignment = Alignment(wrap_text=True, vertical="top")
            options = {
                "F": '"无增生,轻度,轻-中度,中度,中-重度,重度,未提及,不明确"',
                "H": '"形成,未见,不明确,未提及"',
                "I": '"无融合,节段性或部分或少数,较广泛或大部分,广泛性,未提及,不明确"',
            }
            for col, formula in options.items():
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{ws.max_row}")
            ws.freeze_panes = "C2"

    manifest = {
        "seed": SEED,
        "per_stratum": PER_STRATUM,
        "n_sample": len(rows),
        "population": {"linked_model_records": int(len(linked)), "by_split": {s: int((linked["split"] == s).sum()) for s in STRATA}},
        "frozen_reproduction_checks": checks,
        "registry_sha256": sha256_file(rvm.REGISTRY),
        "morphometry_sha256": sha256_file(rvm.MORPHOMETRY),
        "modeling_script_sha256": sha256_file(REPO_SCRIPTS / "remote_v2_modeling.py"),
        "sample_composition": {
            "by_split": {s: sum(r["split"] == s for r in rows) for s in STRATA},
            "by_class": pd.Series([r["disease_code"] for r in rows]).value_counts().to_dict(),
            "by_link_method": pd.Series([r["link_method"] for r in rows]).value_counts().to_dict(),
            "parser_nonmissing": {k: sum(r["parsed"][k] is not None for r in rows) for k, _, _ in FIELD_SPECS},
        },
        "outputs": {
            "key": key_path.name,
            "key_sha256": sha256_file(key_path),
            "worksheets": [f"audit_worksheet_reader{i}.xlsx" for i in (1, 2)],
        },
        "privacy_note": "读取表只含匿名 PT 报告的电镜观察与诊断段落；密钥与读取表均不得进入公开仓库。",
    }
    (OUT_DIR / "audit_sample_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
