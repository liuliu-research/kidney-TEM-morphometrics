"""严格核验双读者/仲裁证据后聚合；人工编码独立于待测解析器。

必须提供两个独立具名读者；200份、七字段均完整，并由第三人裁定所有分歧。
未提及/不明确是显式答案，空白和非法值是错误。试读不能标为human。
四字段以上不一致仅为描述性标记，不能据此推断连接错误或诊断准确性。
"""

from __future__ import annotations

import argparse
import json
import math
import os
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from audit_contract import (
    FIELDS,
    FORBIDDEN_SOURCE,
    AuditValidationError,
    answer_status,
    clean,
    sha256,
    validate_answer,
    validate_reader_metadata,
    validate_sample_answers,
    validate_source_path,
)

HERE = Path(__file__).resolve().parent
PROJECT_ROOT = Path(os.environ.get("TEM_AUDIT_SOURCE_ROOT", str(HERE.parent))).resolve()
TOLERANCE = {
    "glomerulus_count": 0.0,
    "sclerosis_rate": 0.005,
    "mesangial_cell_count": 0.05,
    "gbm_thickness_nm": 0.0,
}
encode = validate_answer


def is_missing(value: object) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value))


def same(field: str, a: float, b: float) -> bool:
    """机器-人工比较使用公开容差；读者分歧判定不用这些容差。"""
    if is_missing(a) and is_missing(b):
        return True
    if is_missing(a) or is_missing(b):
        return False
    tol = TOLERANCE.get(field, 0.0)
    if field == "gbm_thickness_nm":
        tol = max(10.0, 0.01 * abs(b))
    return abs(a - b) <= tol


def wilson(k: int, n: int, z: float = 1.959964) -> tuple[float, float]:
    p = k / n
    denom = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / denom
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denom
    return max(0.0, centre - half), min(1.0, centre + half)


def kappa(a: list[str], b: list[str]) -> float | None:
    n = len(a)
    cats = sorted(set(a) | set(b))
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return None if pe == 1 else (po - pe) / (1 - pe)


def validate_workbook_source(
    path: Path, role: str, name: str, key_hash: str, source_type: str
) -> None:
    """核对工作簿内来源；重命名试读文件不足以升级为人工证据。"""
    validate_source_path(path, source_type)
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if "审计来源" not in book.sheetnames:
            raise AuditValidationError(f"{path.name} 缺少 审计来源 工作表")
        cells = list(book["审计来源"].iter_rows(values_only=True))
        if not cells or tuple(cells[0][:2]) != ("字段", "值"):
            raise AuditValidationError("审计来源 必须使用 字段/值 两列")
        metadata: dict[str, str] = {}
        for row in cells[1:]:
            if not row or not clean(row[0]):
                continue
            key = clean(row[0])
            if key in metadata:
                raise AuditValidationError(f"来源字段重复：{key}")
            metadata[key] = clean(row[1] if len(row) > 1 else "")
        expected = {
            "schema_version": "1",
            "source_type": source_type,
            "role": role,
            "reader_name": name,
            "audit_key_sha256": key_hash,
        }
        if any(metadata.get(key) != val for key, val in expected.items()):
            raise AuditValidationError(f"{path.name} 来源、身份、角色或抽样版本不一致")
        markers = [*book.sheetnames, *metadata.values()]
        markers += [
            clean(getattr(book.properties, key, None))
            for key in ("title", "subject", "description", "keywords")
        ]
        if source_type == "human" and FORBIDDEN_SOURCE.search(" ".join(markers)):
            raise AuditValidationError(f"{path.name} 含试读/合成工作簿标记")
    finally:
        book.close()


def read_sheet(path: Path, expected_ids: set[str]) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="读取表", dtype=str, keep_default_na=False)
    df.columns = [clean(c) for c in df.columns]
    required = {"样本号", *(field[1] for field in FIELDS)}
    if not required.issubset(df.columns):
        raise AuditValidationError(
            f"{path.name} 缺少列：{sorted(required - set(df.columns))}"
        )
    ids = [clean(value) for value in df["样本号"]]
    if any(not value for value in ids) or len(ids) != len(set(ids)):
        raise AuditValidationError(f"{path.name} 含空白或重复样本号")
    if set(ids) != expected_ids:
        raise AuditValidationError(
            f"{path.name} ID不闭合：缺少{len(expected_ids - set(ids))}，"
            f"多余{len(set(ids) - expected_ids)}"
        )
    df["样本号"] = ids
    return df.set_index("样本号")


def raw_sheet(df: pd.DataFrame) -> dict[str, dict[str, str]]:
    raw = {
        str(sid): {key: clean(row[col]) for key, col, *_ in FIELDS}
        for sid, row in df.iterrows()
    }
    for sid, answers in raw.items():
        try:
            validate_sample_answers(answers)
        except AuditValidationError as exc:
            raise AuditValidationError(f"样本{sid}：{exc}") from exc
    return raw


def reader_label(field: str, answers: dict[str, str]) -> str:
    total = answers["glomerulus_count"]
    value = validate_answer(field, answers[field], total)
    status = answer_status(field, answers[field], total)
    if status == "derived_missing_denominator":
        return status + ":" + clean(answers[field])
    return status if is_missing(value) else f"{value:.12g}"


def load_key(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    values = payload.get("rows", [])
    ids = [clean(row.get("sample_id")) for row in values]
    if len(ids) != 200 or len(set(ids)) != 200 or any(not sid for sid in ids):
        raise AuditValidationError("抽样密钥必须为200个不重复非空样本号")
    if Counter(row.get("split") for row in values) != {
        "development": 100,
        "temporal_test": 100,
    }:
        raise AuditValidationError("抽样密钥必须为开发期100份、时间测试期100份")
    if sum("reextracted" in row for row in values) not in {0, 200}:
        raise AuditValidationError("重抽取来源不完整")
    for row in values:
        for attr in ("parsed", "reextracted"):
            if attr not in row and attr == "reextracted":
                continue
            if not isinstance(row.get(attr), dict):
                raise AuditValidationError(f"机器来源缺少{attr}")
            for field, *_ in FIELDS:
                if field not in row[attr]:
                    raise AuditValidationError(f"机器来源缺少字段{field}")
                value = row[attr][field]
                if value is not None and (
                    not isinstance(value, (float, int)) or not math.isfinite(value)
                ):
                    raise AuditValidationError(f"机器来源含非法非有限数值{field}")
    return dict(zip(ids, values))


def resolve_readers(
    r1: dict,
    r2: dict,
    adjudication: Path | None,
    metadata: dict,
    key_hash: str,
    source_type: str,
    todo_path: Path,
    write: bool = True,
) -> tuple[dict, dict, dict]:
    """闭合全部分歧，保留工作簿与声明哈希以供追溯。"""
    differences = {
        (sid, field)
        for sid in r1
        for field, *_ in FIELDS
        if reader_label(field, r1[sid]) != reader_label(field, r2[sid])
    }
    decisions: dict[tuple[str, str], dict[str, str]] = {}
    if adjudication is not None:
        validate_workbook_source(
            adjudication,
            "adjudicator",
            metadata["adjudicator"]["name"],
            key_hash,
            source_type,
        )
        if metadata.get("adjudication_sha256") != sha256(adjudication):
            raise AuditValidationError("仲裁声明未绑定当前工作簿")
        frame = pd.read_excel(
            adjudication, sheet_name="裁定表", dtype=str, keep_default_na=False
        )
        needed = {"样本号", "字段", "读者1", "读者2", "裁定值", "裁定依据"}
        if not needed.issubset(frame.columns):
            raise AuditValidationError("仲裁表缺少必需列或裁定依据列")
        for _, row in frame.iterrows():
            pair = clean(row["样本号"]), clean(row["字段"])
            if pair in decisions or pair not in differences:
                raise AuditValidationError("仲裁表含重复、额外或非分歧样本-字段对")
            sid, field = pair
            if (
                clean(row["读者1"]) != r1[sid][field]
                or clean(row["读者2"]) != r2[sid][field]
            ):
                raise AuditValidationError("仲裁表原始读者答案与当前版本不一致")
            if not clean(row["裁定值"]) or not clean(row["裁定依据"]):
                raise AuditValidationError("仲裁值及依据不得留空")
            decisions[pair] = {clean(key): clean(value) for key, value in row.items()}
        if set(decisions) != differences:
            raise AuditValidationError("仲裁表未完整覆盖全部读者分歧")
    elif differences:
        if write:
            todo = [
                {
                    "样本号": sid,
                    "字段": field,
                    "读者1": r1[sid][field],
                    "读者2": r2[sid][field],
                    "裁定值": "",
                    "肾小球总数（供硬化率换算）": "",
                    "裁定依据": "",
                }
                for sid, field in sorted(differences)
            ]
            pd.DataFrame(todo).to_excel(todo_path, sheet_name="裁定表", index=False)
        raise AuditValidationError(f"{len(differences)}处分歧待具名仲裁")
    resolved = {}
    raw_final: dict[str, dict[str, str]] = {}
    status_counts = {field: Counter() for field, *_ in FIELDS}
    for sid in r1:
        raw = {
            field: (
                decisions[(sid, field)]["裁定值"]
                if (sid, field) in decisions
                else r1[sid][field]
            )
            for field, *_ in FIELDS
        }
        raw_final[sid] = dict(raw)
        resolved[sid] = validate_sample_answers(raw)
        for field, *_ in FIELDS:
            status_counts[field][
                answer_status(field, raw[field], raw["glomerulus_count"])
            ] += 1
        pair = sid, "sclerosis_rate"
        if pair in decisions:
            denominator = decisions[pair].get("肾小球总数（供硬化率换算）", "")
            if denominator and not same(
                "glomerulus_count",
                validate_answer("glomerulus_count", denominator),
                resolved[sid]["glomerulus_count"],
            ):
                raise AuditValidationError("仲裁辅助分母与最终已裁定总数不一致")
    agreement = {}
    for field, _, _, kind in FIELDS:
        a = [reader_label(field, r1[sid]) for sid in r1]
        b = [reader_label(field, r2[sid]) for sid in r1]
        k = kappa(a, b) if kind != "numeric" else None
        agreement[field] = {
            "n": len(r1),
            "percent_agreement": round(
                100 * sum(x == y for x, y in zip(a, b)) / len(r1), 1
            ),
            "kappa": round(k, 3) if k is not None else None,
        }
    return (
        resolved,
        agreement,
        {
            "n_disagreements": len(differences),
            "n_resolved": len(decisions),
            "complete": True,
            "adjudicator": metadata["adjudicator"]["name"],
            "final_answer_status_counts": {
                key: dict(counts) for key, counts in status_counts.items()
            },
            "source": str(adjudication.resolve()) if adjudication else None,
            "sha256": sha256(adjudication) if adjudication else None,
            "raw_final": raw_final,
        },
    )


def apply_rereview(
    path: Path,
    raw_final: dict,
    adjudicated: dict,
    r1: dict,
    r2: dict,
    metadata: dict,
    key_hash: str,
    source_type: str,
) -> tuple[dict, dict]:
    """事后文本一致性复裁：仲裁者只对自动检查标记的格子按原文重裁；原始读者答案与主裁定不改。"""
    validate_source_path(path, source_type)
    frame = pd.read_excel(path, sheet_name="复裁表", dtype=str, keep_default_na=False)
    needed = {"样本号", "字段", "来源", "读者1", "读者2", "原参考值", "复裁值", "复裁依据", "标记原因"}
    if not needed.issubset(frame.columns):
        raise AuditValidationError("复裁表缺少必需列")
    source = pd.read_excel(path, sheet_name="审计来源", dtype=str, keep_default_na=False)
    info = {clean(k): clean(v) for k, v in zip(source["字段"], source["值"])}
    if (
        info.get("role") != "adjudicator"
        or info.get("kind") != "rereview"
        or info.get("audit_key_sha256") != key_hash
        or info.get("source_type") != source_type
        or info.get("reader_name") != clean(metadata["adjudicator"]["name"])
    ):
        raise AuditValidationError("复裁表来源声明与当前仲裁者或审计版本不一致")
    cells, seen = [], set()
    new_raw = {sid: dict(values) for sid, values in raw_final.items()}
    for _, row in frame.iterrows():
        sid, field = clean(row["样本号"]), clean(row["字段"])
        if (sid, field) in seen or sid not in raw_final or field not in {key for key, *_ in FIELDS}:
            raise AuditValidationError("复裁表含重复或无效的样本-字段对")
        seen.add((sid, field))
        if (
            clean(row["读者1"]) != r1[sid][field]
            or clean(row["读者2"]) != r2[sid][field]
            or clean(row["原参考值"]) != raw_final[sid][field]
        ):
            raise AuditValidationError("复裁表原始答案与当前版本不一致")
        value, reason = clean(row["复裁值"]), clean(row["复裁依据"])
        if not value or not reason:
            raise AuditValidationError("复裁值及依据不得留空")
        new_raw[sid][field] = value
        cells.append(
            {
                "sample_id": sid,
                "field": field,
                "origin": clean(row["来源"]),
                "reader1": r1[sid][field],
                "reader2": r2[sid][field],
                "before": raw_final[sid][field],
                "after": value,
                "changed": value != raw_final[sid][field],
                "flag_reason": clean(row["标记原因"]),
                "reason": reason,
            }
        )
    resolved = dict(adjudicated)
    status_after = {field: Counter() for field, *_ in FIELDS}
    for sid, raw in new_raw.items():
        resolved[sid] = validate_sample_answers(raw)
        for field, *_ in FIELDS:
            status_after[field][answer_status(field, raw[field], raw["glomerulus_count"])] += 1
    return resolved, {
        "kind": "post_hoc_text_consistency_rereview",
        "n_cells": len(cells),
        "n_changed": sum(cell["changed"] for cell in cells),
        "fields_affected": sorted({cell["field"] for cell in cells}),
        "cells": cells,
        "adjudicator": metadata["adjudicator"]["name"],
        "source": str(path.resolve()),
        "sha256": sha256(path),
        "queue_sha256": info.get("queue_sha256"),
        "final_answer_status_counts_after": {k: dict(v) for k, v in status_after.items()},
    }


def aggregate_results(rows: dict, adjudicated: dict) -> dict:
    """保持稿件消费的聚合字段；标记只描述不一致。"""
    pipelines = {"frozen_pipeline": "parsed"}
    if all("reextracted" in row for row in rows.values()):
        pipelines["reextracted_pipeline"] = "reextracted"
    result = {}
    for pname, attr in pipelines.items():
        fields_out = {}
        record_mismatch = Counter()
        for key_name, _, eng, kind in FIELDS:
            per_split = {}
            for split in ("all", "development", "temporal_test"):
                c = Counter()
                for sid, row in rows.items():
                    if split != "all" and row["split"] != split:
                        continue
                    value = row[attr][key_name]
                    machine = float("nan") if value is None else float(value)
                    human = adjudicated[sid][key_name]
                    if is_missing(machine) and is_missing(human):
                        c["agree_both_missing"] += 1
                    elif is_missing(machine):
                        c["parser_missing_reader_value"] += 1
                    elif is_missing(human):
                        c["parser_value_reader_missing"] += 1
                    elif same(key_name, machine, human):
                        c["agree_both_value"] += 1
                    else:
                        c["value_differs"] += 1
                    if split == "all" and not same(key_name, machine, human):
                        record_mismatch[sid] += 1
                n = sum(c.values())
                err = sum(
                    c[key]
                    for key in (
                        "value_differs",
                        "parser_missing_reader_value",
                        "parser_value_reader_missing",
                    )
                )
                lo, hi = wilson(err, n)
                both = c["agree_both_value"] + c["value_differs"]
                per_split[split] = {
                    "n": n,
                    "parser_value_n": both + c["parser_value_reader_missing"],
                    "reader_value_n": both + c["parser_missing_reader_value"],
                    **{
                        key: c[key]
                        for key in (
                            "agree_both_value",
                            "agree_both_missing",
                            "value_differs",
                            "parser_missing_reader_value",
                            "parser_value_reader_missing",
                        )
                    },
                    "error_n": err,
                    "error_pct": round(100 * err / n, 1),
                    "error_ci_low_pct": round(100 * lo, 1),
                    "error_ci_high_pct": round(100 * hi, 1),
                    "value_disagreement_pct_when_both_present": (
                        round(100 * c["value_differs"] / both, 1) if both else None
                    ),
                }
            fields_out[key_name] = {"label": eng, "kind": kind, **per_split}
        flagged = sorted(sid for sid, count in record_mismatch.items() if count >= 4)
        flags = {
            "threshold_fields": 4,
            "n_records": len(flagged),
            "sample_ids": flagged,
            "link_methods": dict(
                Counter(rows[sid].get("link_method", "not_recorded") for sid in flagged)
            ),
            "interpretation": (
                "Descriptive multiple-field disagreement only; "
                "not evidence of linkage error."
            ),
        }
        errors = [fields_out[key]["all"]["error_pct"] for key, *_ in FIELDS]
        result[pname] = {
            "fields": fields_out,
            "error_pct_range": [min(errors), max(errors)],
            "disagreement_flags": flags,
            "linkage_flags": flags,
        }
    return result


def run(args: argparse.Namespace, *, write: bool = True) -> dict:
    """全部准入通过后才输出；失败不覆盖旧结果。write=False仅做只读复验。"""
    source_type = args.label
    source_paths = [args.reader1, args.reader2, args.readers, args.key]
    if args.adjudication is not None:
        source_paths.append(args.adjudication)
    rereview_path = getattr(args, "rereview", None)
    if rereview_path is not None:
        source_paths.append(rereview_path)
    for path in source_paths:
        if not path.is_file():
            raise AuditValidationError(f"缺少必需证据：{path.name}")
        validate_source_path(path, source_type)
        if source_type == "human" and not path.resolve().is_relative_to(PROJECT_ROOT):
            raise AuditValidationError("人工证据必须保存在本项目目录内")
    if args.reader1.resolve() == args.reader2.resolve():
        raise AuditValidationError("两位读者不能使用同一个工作簿")
    if source_type != "human" and args.out.name == "audit_results.json":
        raise AuditValidationError("合成测试禁止写入正式audit_results.json")
    rows = load_key(args.key)
    key_hash = sha256(args.key)
    sample_manifest = HERE / "audit_sample_manifest.json"
    if source_type == "human":
        sample = json.loads(sample_manifest.read_text(encoding="utf-8"))
        if (
            sample.get("n_sample") != 200
            or sample.get("outputs", {}).get("key_sha256") != key_hash
        ):
            raise AuditValidationError("人工核对必须使用冻结抽样清单绑定的200份样本")
    metadata = json.loads(args.readers.read_text(encoding="utf-8"))
    validate_reader_metadata(metadata, key_hash, source_type)
    raw = []
    for role, path in (("reader1", args.reader1), ("reader2", args.reader2)):
        validate_workbook_source(
            path, role, metadata[role]["name"], key_hash, source_type
        )
        if metadata[role].get("worksheet_sha256") != sha256(path):
            raise AuditValidationError(f"{role}声明未绑定当前工作簿")
        raw.append(raw_sheet(read_sheet(path, set(rows))))
    adjudicated, agreement, adjudication = resolve_readers(
        raw[0],
        raw[1],
        args.adjudication,
        metadata,
        key_hash,
        source_type,
        args.out.parent / "audit_adjudication_todo.xlsx",
        write,
    )
    raw_final = adjudication.pop("raw_final")
    rereview = None
    if rereview_path is not None:
        if args.adjudication is None:
            raise AuditValidationError("复裁只能在主裁定完成后进行")
        before = aggregate_results(rows, adjudicated)
        adjudicated, rereview = apply_rereview(
            rereview_path, raw_final, adjudicated, raw[0], raw[1], metadata, key_hash, source_type
        )
        rereview["pipelines_before_rereview"] = {
            name: {
                "error_pct_range": pipeline["error_pct_range"],
                "disagreement_flags": pipeline["disagreement_flags"],
                "fields": {f: pipeline["fields"][f] for f in rereview["fields_affected"]},
            }
            for name, pipeline in before.items()
        }
    results = {
        "label": source_type,
        "status": "VERIFIED" if source_type == "human" else "SYNTHETIC_TEST_ONLY",
        "n": len(rows),
        "by_split": dict(Counter(row["split"] for row in rows.values())),
        "readers": metadata,
        "dual_reader": True,
        "inter_reader_before_adjudication": agreement,
        "adjudication": adjudication,
        "rereview": rereview,
        "validation": {
            "contract_version": 1,
            "status": "PASS",
            "exact_sample_ids": True,
            "all_seven_fields_complete": True,
            "independent_reference_encoder": True,
            "identity_scope": (
                "Named participant declarations; real-world identity is "
                "not independently verified by this script."
            ),
            "limitations": [
                "Missing/ambiguous report text is not a verified clinical negative.",
                "Field agreement is not diagnostic accuracy or proof of linkage correctness.",
                "This service/sample audit is descriptive, not external validation.",
            ],
            "machine_comparison_tolerance": {
                **TOLERANCE,
                "gbm_thickness_nm": "max(10 nm, 1% of human value)",
            },
            "range_coding": "Intervals use midpoint; GBM units are converted to nm.",
            "reader_agreement": (
                "Exact encoded value and explicit missingness state; "
                "machine comparison tolerances are not applied."
            ),
        },
        "input_provenance": {
            "key_sha256": key_hash,
            "sample_manifest_sha256": (
                sha256(sample_manifest) if source_type == "human" else None
            ),
            "reader1_sha256": sha256(args.reader1),
            "reader2_sha256": sha256(args.reader2),
            "readers_sha256": sha256(args.readers),
            "contract_sha256": sha256(HERE / "audit_contract.py"),
            "comparison_script_sha256": sha256(Path(__file__)),
            "paths": {
                role: str(path.resolve())
                for role, path in (
                    ("key", args.key),
                    ("reader1", args.reader1),
                    ("reader2", args.reader2),
                    ("readers", args.readers),
                )
                + ((("rereview", rereview_path),) if rereview_path is not None else ())
            },
        },
        "pipelines": aggregate_results(rows, adjudicated),
    }
    if write:
        args.out.write_text(
            json.dumps(results, ensure_ascii=False, indent=2, allow_nan=False),
            encoding="utf-8",
        )
    return results


def validate_verified_audit_result(path: Path) -> dict:
    """供稿件/回复构建入口调用：只读重算来源并核对结果，不接受旧pilot或漂移。"""
    path = Path(path)
    try:
        claimed = json.loads(path.read_text(encoding="utf-8"))
        if claimed.get("label") != "human" or claimed.get("status") != "VERIFIED":
            raise AuditValidationError("稿件只接收已核验的真实人工结果")
        paths = claimed["input_provenance"]["paths"]
        adjudication = claimed["adjudication"]["source"]
        args = argparse.Namespace(
            **{
                key: Path(paths[key])
                for key in ("key", "reader1", "reader2", "readers")
            },
            adjudication=Path(adjudication) if adjudication else None,
            rereview=Path(paths["rereview"]) if paths.get("rereview") else None,
            out=path,
            label="human",
        )
        actual = run(args, write=False)
        if claimed != actual:
            raise AuditValidationError("人工结果或来源/工具版本漂移；须重跑人工比对")
        return actual
    except (OSError, KeyError, TypeError, ValueError) as exc:
        raise AuditValidationError(f"人工结果复验失败：{exc}") from exc


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reader1", type=Path, required=True)
    parser.add_argument("--reader2", type=Path, required=True)
    parser.add_argument("--adjudication", type=Path)
    parser.add_argument("--rereview", type=Path, help="事后文本一致性复裁表（门户 --export-rereview 导出）")
    parser.add_argument("--readers", type=Path, default=HERE / "audit_readers.json")
    parser.add_argument(
        "--key", type=Path, default=HERE / "audit_key_CONFIDENTIAL.json"
    )
    parser.add_argument("--out", type=Path, default=HERE / "audit_results.json")
    parser.add_argument("--label", choices=("human", "synthetic"), default="human")
    args = parser.parse_args()
    try:
        results = run(args)
    except (AuditValidationError, OSError, ValueError, KeyError) as exc:
        raise SystemExit(f"[BLOCKED] {exc}") from exc
    print(
        json.dumps(
            {
                "status": results["status"],
                "label": results["label"],
                "n": results["n"],
                "dual_reader": True,
                "adjudicated": results["adjudication"]["n_resolved"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
